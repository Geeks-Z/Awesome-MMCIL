#!/usr/bin/env python3
"""Write all complete OpenAI CLIP ENGINE and BOFA logs into their seed sheets."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from reconcile_clip_workbooks import OPENAI_LOG_ROOT, close_pair, direct_results, workbook_locations


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenAI_CLIP_ViT-B16_Baselines.xlsx"
METHODS = ("ENGINE", "BOFA")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="write updates to the workbook")
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    results = direct_results(OPENAI_LOG_ROOT, METHODS, "openai_clip")
    workbook = load_workbook(args.workbook)
    inserted = already_present = 0

    for (seed, method, dataset, initial, increment), metrics in sorted(results.items()):
        sheet = f"Seed{seed}"
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet: {sheet}")
        worksheet = workbook[sheet]
        row, column = workbook_locations(worksheet)[(method, (dataset, initial, increment))]
        current = (worksheet.cell(row, column).value, worksheet.cell(row, column + 1).value)
        rounded = (round(metrics[0], 2), round(metrics[1], 2))

        if current == (None, None):
            worksheet.cell(row, column).value = rounded[0]
            worksheet.cell(row, column + 1).value = rounded[1]
            print(f"insert Seed{seed} {method} {dataset} B{initial} Inc{increment}: {rounded}")
            inserted += 1
        elif all(isinstance(value, (int, float)) for value in current) and close_pair(
            (float(current[0]), float(current[1])), metrics
        ):
            already_present += 1
        else:
            raise ValueError(
                f"Refusing to overwrite {sheet} {method} {dataset} B{initial} Inc{increment}: "
                f"workbook={current}, log={rounded}"
            )

    print(f"complete_logs={len(results)} inserted={inserted} already_present={already_present}")
    if args.apply:
        workbook.save(args.workbook)
        print(f"Updated: {args.workbook}")
    else:
        print("Dry run only; pass --apply to save.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
