#!/usr/bin/env python3
"""Extract completed AREA runs from logs and update MMCL_Baselines.xlsx."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_Baselines.xlsx"
LOG_DIRS = (ROOT / "logs" / "AREA", ROOT / "logs" / "area", ROOT / "logs" / "aera")

SEED_SHEETS = {1993: "Seed1993", 2026: "Seed2026", 0: "Seed0", 42: "Seed42"}

AREA_BLOCKS = (
    ("aircraft_clip_0_10", "aircraft_clip_50_10", "cifar224_clip_0_10", "cifar224_clip_50_10", "cars_clip_0_10", "cars_clip_50_10"),
    ("imagenetr_clip_0_20", "imagenetr_clip_100_20", "cub_clip_0_20", "cub_clip_100_20", "ucf101_clip_0_10", "ucf101_clip_50_10"),
    ("sun_clip_0_30", "sun_clip_150_30", "food101_clip_0_10", "food101_clip_50_10", "objectnet_clip_0_20", "objectnet_clip_100_20"),
)


def log_stem(config: Path) -> str:
    args = json.loads(config.read_text())
    initial = 0 if args["init_cls"] == args["increment"] else args["init_cls"]
    return f"{args['dataset']}_{args['backbone_type']}_{initial}_{args['increment']}"


def completed_metrics(log_path: Path) -> tuple[float, float] | None:
    """Return (A_bar, A_B) only when the final task has completed."""
    if not log_path.is_file():
        return None

    content = log_path.read_text(errors="replace").rstrip()
    if not content:
        return None

    averages = re.findall(r"Average Accuracy \(CNN top1\): ([0-9.]+)", content)
    curves = re.findall(r"CNN top1 curve: \[([^]]+)\]", content)
    if not averages or not curves or not content.splitlines()[-1].endswith(averages[-1]):
        return None

    curve = [float(value.strip()) for value in curves[-1].split(",")]
    return round(float(averages[-1]), 2), round(curve[-1], 2)


def area_log_path(stem: str, seed: int) -> Path:
    """Prefer the canonical AREA directory and read older logs if necessary."""
    for directory in LOG_DIRS:
        path = directory / f"{stem}_{seed}.log"
        if path.is_file():
            return path
    return LOG_DIRS[0] / f"{stem}_{seed}.log"


def is_complete(config: Path, seed: int) -> bool:
    return completed_metrics(area_log_path(log_stem(config), seed)) is not None


def update_workbook() -> list[str]:
    workbook = load_workbook(WORKBOOK)
    updates: list[str] = []

    for seed, sheet_name in SEED_SHEETS.items():
        sheet = workbook[sheet_name]
        headers = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row=row, column=1).value == "Method"]
        if len(headers) != len(AREA_BLOCKS):
            raise ValueError(f"{sheet_name}: expected {len(AREA_BLOCKS)} Method headers, found {len(headers)}")
        for header, next_header, stems in zip(headers, (*headers[1:], sheet.max_row + 1), AREA_BLOCKS):
            row = next(
                (
                    candidate
                    for candidate in range(header + 2, next_header)
                    if sheet.cell(row=candidate, column=1).value == "AREA"
                ),
                None,
            )
            # AREA may have been deliberately removed from a block; do not
            # insert or overwrite another method merely to create it.
            if row is None:
                continue
            sheet.cell(row=row, column=1).value = "AREA"
            sheet.cell(row=row, column=2).value = "ICML 2026"
            for column in range(3, 15):
                sheet.cell(row=row, column=column).value = None
            for index, stem in enumerate(stems):
                metrics = completed_metrics(area_log_path(stem, seed))
                if metrics is None:
                    continue
                a_bar, a_b = metrics
                a_bar_column = 3 + index * 2
                a_b_column = a_bar_column + 1
                sheet.cell(row=row, column=a_bar_column).value = a_bar
                sheet.cell(row=row, column=a_b_column).value = a_b
                updates.append(
                    f"{sheet_name}!{sheet.cell(row=row, column=a_bar_column).coordinate}:{a_bar}, "
                    f"{sheet.cell(row=row, column=a_b_column).coordinate}:{a_b}"
                )

    workbook.save(WORKBOOK)
    return updates


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--is-complete",
        nargs=2,
        metavar=("CONFIG", "SEED"),
        help="return success when this config/seed has a completed AREA log",
    )
    args = parser.parse_args()

    if args.is_complete:
        config, seed = args.is_complete
        return 0 if is_complete(ROOT / config, int(seed)) else 1

    updates = update_workbook()
    print("AREA workbook update complete")
    for update in updates:
        print(update)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
