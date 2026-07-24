#!/usr/bin/env python3
"""Update audited baseline results in MMCL_Baselines.xlsx by method and header.

The workbook layout is intentionally discovered from each ``Method`` / ``Venue`` /
``Year`` header.  No result row or metric column is addressed by a fixed number.
Only complete results are written; unmatched protocols are cleared for the audited
method on the target seed so stale values cannot be mislabelled as a valid run.
"""

from __future__ import annotations

import argparse
import ast
import math
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_Baselines.xlsx"
TARGET_SEED = 1993
TARGET_SHEET = "Seed1993"

METHOD_VENUES: dict[str, tuple[str, int]] = {
    "SimpleCIL": ("IJCV", 2024),
    "ZS-CLIP": ("ICML", 2021),
    "L2P": ("CVPR", 2022),
    "CODA-Prompt": ("CVPR", 2023),
    "DualPrompt": ("ECCV", 2022),
    "RAPF": ("ECCV", 2024),
    "PROOF": ("TPAMI", 2025),
    "ENGINE": ("ICCV", 2025),
    "CLG-CBM": ("CVPR", 2025),
    "BOFA": ("AAAI", 2026),
    "PromptFusion": ("ECCV", 2024),
    "MoE-Adapters": ("CVPR", 2024),
    "AREA": ("ICML", 2026),
}

RESULT_FILES = {
    "L2P": ROOT / "results" / "L2P-LAION-400M-SEED_1993-A40.out",
    "CODA-Prompt": ROOT / "results" / "CODA-Prompt-LAION-400M-SEED_1993-A40.out",
    "DualPrompt": ROOT / "results" / "DualPrompt-LAION-400M-SEED_1993-A40.out",
    "PROOF": ROOT / "results" / "PROOF-LAION-400M-SEED_1993-A40.out",
}

DATASET_ALIASES = {
    "aircraft": "aircraft",
    "cifar100": "cifar224",
    "cifar224": "cifar224",
    "cars": "cars",
    "inr": "imagenetr",
    "imagenetr": "imagenetr",
    "cub": "cub",
    "ucf": "ucf101",
    "sun": "sun",
    "food": "food101",
    "objectnet": "objectnet",
}

DATASET_CLASS_COUNTS = {
    "aircraft": 100,
    "cifar224": 100,
    "cars": 196,
    "imagenetr": 200,
    "cub": 200,
    "ucf101": 101,
    "sun": 300,
    "food101": 101,
    "objectnet": 200,
}


def _expected_task_count(dataset: str, init_cls: int, increment: int) -> int:
    """Return the number of incremental tasks for a supported protocol."""
    total_classes = DATASET_CLASS_COUNTS[dataset]
    return 1 + math.ceil((total_classes - init_cls) / increment)


def _normalise_dataset(dataset: str) -> str | None:
    return DATASET_ALIASES.get(dataset.lower().replace("-", "").replace("_", ""))


def _protocol_from_log_filename(path: Path) -> tuple[str, int, int] | None:
    """Read dataset/Base/Inc from standard per-run log filenames."""
    match = re.fullmatch(r"(.+?)_(?:openai_clip|vit_b16|clip)_(\d+)_(\d+)_\d+\.log", path.name)
    if not match:
        return None
    dataset_name, base, increment = match.groups()
    dataset = _normalise_dataset(dataset_name)
    if dataset is None:
        return None
    base, increment = int(base), int(increment)
    return dataset, increment if base == 0 else base, increment


def _structured_log_result(
    block: str, filename_protocol: tuple[str, int, int] | None
) -> tuple[tuple[str, int, int], tuple[float, float]] | None:
    """Read a complete FileHandler log that has no terminal ``Finished`` line."""
    dataset_matches = re.findall(r"\[trainer\.py\] => dataset:\s*(\S+)", block)
    init_matches = re.findall(r"\[trainer\.py\] => init_cls:\s*(\d+)", block)
    increment_matches = re.findall(r"\[trainer\.py\] => increment:\s*(\d+)", block)
    average_matches = re.findall(r"Average Accuracy \(CNN top1\):\s*([0-9.]+)", block)
    curve_matches = re.findall(r"CNN top1 curve:\s*(\[[^\n]+\])", block)
    if not (average_matches and curve_matches):
        return None
    if dataset_matches and init_matches and increment_matches:
        dataset = _normalise_dataset(dataset_matches[-1])
        init_cls, increment = int(init_matches[-1]), int(increment_matches[-1])
    elif filename_protocol is not None:
        dataset, init_cls, increment = filename_protocol
    else:
        return None
    if dataset not in DATASET_CLASS_COUNTS:
        return None
    try:
        curve = ast.literal_eval(curve_matches[-1])
    except (SyntaxError, ValueError):
        return None
    if not isinstance(curve, list) or len(curve) != _expected_task_count(dataset, init_cls, increment):
        return None
    if not all(isinstance(value, (int, float)) for value in curve):
        return None
    return (dataset, init_cls, increment), (
        round(float(average_matches[-1]), 2),
        round(float(curve[-1]), 2),
    )


def completed_results(path: Path) -> dict[tuple[str, int, int], tuple[float, float]]:
    """Extract one validated (A_bar, A_B) pair from each complete config block."""
    content = path.read_text(errors="replace")
    starts = [match.start() for match in re.finditer(r"\[trainer\.py\] => config:", content)]
    if not starts:
        starts = [0]
    starts.append(len(content))
    filename_protocol = _protocol_from_log_filename(path)
    results: dict[tuple[str, int, int], tuple[float, float]] = {}
    for start, end in zip(starts, starts[1:]):
        block = content[start:end]
        protocol = re.search(r"Finished\s+(\S+)_init(\d+)_inc(\d+)", block)
        averages = re.findall(r"Average Accuracy \(Top1\):\s*([0-9.]+)", block)
        lasts = re.findall(r"Last Accuracy:\s*([0-9.]+)", block)
        if protocol and averages and lasts:
            dataset = _normalise_dataset(protocol.group(1))
            if dataset is None:
                continue
            key = (dataset, int(protocol.group(2)), int(protocol.group(3)))
            metrics = (round(float(averages[-1]), 2), round(float(lasts[-1]), 2))
        else:
            structured = _structured_log_result(block, filename_protocol)
            if structured is None:
                continue
            key, metrics = structured
        if key in results:
            raise ValueError(f"{path.name}: duplicate completed result for {key}")
        results[key] = metrics
    return results


def headers(ws) -> list[int]:
    return [
        row
        for row in range(1, ws.max_row + 1)
        if ws.cell(row=row, column=1).value == "Method"
    ]


def method_rows(ws, header_row: int, next_header: int) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row in range(header_row + 1, next_header):
        method = ws.cell(row=row, column=1).value
        if isinstance(method, str) and method.strip():
            if method in rows:
                raise ValueError(f"{ws.title}: duplicate {method!r} below row {header_row}")
            rows[method] = row
    return rows


def protocol_columns(ws, header_row: int) -> list[tuple[int, tuple[str, int, int]]]:
    result: list[tuple[int, tuple[str, int, int]]] = []
    for column in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=column).value
        if not isinstance(value, str):
            continue
        match = re.fullmatch(r"(.+?) B(\d+) Inc(\d+)", value.strip())
        if not match:
            continue
        dataset_label, base, increment = match.groups()
        dataset_key = DATASET_ALIASES.get(dataset_label.lower().replace("-", ""))
        if dataset_key is None:
            raise ValueError(f"{ws.title}!{ws.cell(header_row, column).coordinate}: unsupported dataset {dataset_label!r}")
        init = int(increment) if int(base) == 0 else int(base)
        result.append((column, (dataset_key, init, int(increment))))
    if not result:
        raise ValueError(f"{ws.title}!{header_row}: no protocol headers found")
    return result


def set_cell(ws, row: int, column: int, value, changes: list[str]) -> None:
    cell = ws.cell(row=row, column=column)
    if cell.value != value:
        changes.append(f"{ws.title}!{cell.coordinate}: {cell.value!r} -> {value!r}")
        cell.value = value


def update_venues(workbook, changes: list[str]) -> None:
    for ws in workbook.worksheets:
        if ws.title == "Experiment Status":
            continue
        found_headers = headers(ws)
        if len(found_headers) != 3:
            raise ValueError(f"{ws.title}: expected three Method headers, found {len(found_headers)}")
        for header_row, next_header in zip(found_headers, (*found_headers[1:], ws.max_row + 1)):
            if ws.cell(row=header_row, column=2).value != "Venue" or ws.cell(row=header_row, column=3).value != "Year":
                raise ValueError(f"{ws.title}!{header_row}: expected Venue and Year columns")
            for method, row in method_rows(ws, header_row, next_header).items():
                if method not in METHOD_VENUES:
                    continue
                venue, year = METHOD_VENUES[method]
                set_cell(ws, row, 2, venue, changes)
                set_cell(ws, row, 3, year, changes)


def update_audited_results(workbook, changes: list[str]) -> None:
    ws = workbook[TARGET_SHEET]
    found_headers = headers(ws)
    if len(found_headers) != 3:
        raise ValueError(f"{ws.title}: expected three Method headers, found {len(found_headers)}")
    parsed = {method: completed_results(path) for method, path in RESULT_FILES.items()}
    for header_row, next_header in zip(found_headers, (*found_headers[1:], ws.max_row + 1)):
        rows = method_rows(ws, header_row, next_header)
        for method, results in parsed.items():
            row = rows.get(method)
            if row is None:
                raise ValueError(f"{ws.title}: no {method} row below header {header_row}")
            for a_bar_column, protocol in protocol_columns(ws, header_row):
                metrics = results.get(protocol)
                set_cell(ws, row, a_bar_column, None if metrics is None else metrics[0], changes)
                set_cell(ws, row, a_bar_column + 1, None if metrics is None else metrics[1], changes)


def update_completed_log(
    workbook, method: str, seed: int, result_log: Path, changes: list[str]
) -> None:
    """Write only protocols completed in one standalone log.

    Method rows and result columns are always resolved from their labels, so a
    user may freely insert, remove, or reorder rows in the workbook.
    """
    sheet_name = f"Seed{seed}"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"workbook has no {sheet_name!r} sheet")

    results = completed_results(result_log)
    if not results:
        raise ValueError(f"{result_log}: no complete result block found")

    ws = workbook[sheet_name]
    found_headers = headers(ws)
    if len(found_headers) != 3:
        raise ValueError(f"{ws.title}: expected three Method headers, found {len(found_headers)}")
    matched_protocols: set[tuple[str, int, int]] = set()
    matched_method = False
    for header_row, next_header in zip(found_headers, (*found_headers[1:], ws.max_row + 1)):
        row = method_rows(ws, header_row, next_header).get(method)
        if row is None:
            continue
        matched_method = True
        for a_bar_column, protocol in protocol_columns(ws, header_row):
            metrics = results.get(protocol)
            if metrics is None:
                continue
            matched_protocols.add(protocol)
            set_cell(ws, row, a_bar_column, metrics[0], changes)
            set_cell(ws, row, a_bar_column + 1, metrics[1], changes)

    if not matched_method:
        raise ValueError(f"{ws.title}: no {method!r} row")
    unmatched = sorted(set(results).difference(matched_protocols))
    if unmatched:
        raise ValueError(f"{result_log.name}: no matching workbook protocol for {unmatched}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--method", help="method label to update from one completed log")
    parser.add_argument("--seed", type=int, help="seed corresponding to the workbook sheet")
    parser.add_argument("--result-log", type=Path, help="standalone completed experiment log")
    args = parser.parse_args()
    standalone_args = (args.method, args.seed, args.result_log)
    if any(value is not None for value in standalone_args) and any(
        value is None for value in standalone_args
    ):
        parser.error("--method, --seed, and --result-log must be supplied together")

    workbook = load_workbook(WORKBOOK)
    changes: list[str] = []
    update_venues(workbook, changes)
    if args.result_log is None:
        update_audited_results(workbook, changes)
    else:
        result_log = args.result_log.resolve()
        if not result_log.is_file():
            parser.error(f"result log not found: {result_log}")
        update_completed_log(workbook, args.method, args.seed, result_log, changes)
    print(f"{'Would change' if args.dry_run else 'Changed'} {len(changes)} cells")
    for change in changes:
        print(change)
    if not args.dry_run:
        workbook.save(WORKBOOK)
        print(f"Updated: {WORKBOOK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
