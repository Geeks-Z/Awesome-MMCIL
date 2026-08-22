#!/usr/bin/env python3
"""Print unfilled LAION-400M workbook jobs for a native MMCL method."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenCLIP_ViT-B16_LAION-400M_Baselines.xlsx"
METHOD_DIRECTORIES = {
    "ranpac": "ranpac",
    "fecam": "fecam",
    "area": "area",
    "clg-cbm": "CLG-CBM",
}
METHOD_DISPLAY_NAMES = {
    "ranpac": "RanPAC",
    "fecam": "FeCAM",
    "area": "AREA",
    "clg-cbm": "CLG-CBM",
}
DATASET_LABELS = {
    "aircraft": "Aircraft",
    "cars": "Cars",
    "cifar224": "CIFAR100",
    "cub": "CUB",
    "food101": "Food",
    "imagenetr": "INR",
    "objectnet": "ObjectNet",
    "sun": "SUN",
    "ucf101": "UCF",
}


def canonical_method(name: str) -> str:
    return name.lower().replace("_", "-")


def protocol_for(config: dict) -> str:
    init_cls = int(config["init_cls"])
    increment = int(config["increment"])
    # B0 is represented by init_cls == increment in repository JSON configs.
    display_init = 0 if init_cls == increment else init_cls
    return f"{DATASET_LABELS[config['dataset']]} B{display_init} Inc{increment}"


def find_method_row(sheet, header_row: int, next_header_row: int, method: str) -> int:
    for row in range(header_row + 1, next_header_row):
        if sheet.cell(row, 1).value == method:
            return row
    raise ValueError(f"{method} row not found in {sheet.title}, header row {header_row}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--method", required=True, choices=sorted(METHOD_DIRECTORIES))
    args = parser.parse_args()
    method_key = canonical_method(args.method)
    config_dir = ROOT / "configs" / METHOD_DIRECTORIES[method_key]

    protocol_to_config = {}
    for config_path in sorted(config_dir.glob("*.json")):
        config = json.loads(config_path.read_text())
        if canonical_method(str(config["model_name"])) != method_key:
            raise ValueError(f"Unexpected model_name in {config_path}: {config['model_name']}")
        protocol = protocol_for(config)
        if protocol in protocol_to_config:
            raise ValueError(f"Duplicate protocol config for {protocol}")
        protocol_to_config[protocol] = config_path

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    jobs = []
    for sheet in workbook.worksheets:
        if not sheet.title.startswith("Seed"):
            continue
        seed = int(sheet.title.removeprefix("Seed"))
        header_rows = [
            row
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 1).value == "Method"
        ]
        for index, header_row in enumerate(header_rows):
            next_header = (
                header_rows[index + 1] if index + 1 < len(header_rows) else sheet.max_row + 1
            )
            method_row = find_method_row(
                sheet, header_row, next_header, METHOD_DISPLAY_NAMES[method_key]
            )
            for column in range(4, 16, 2):
                protocol = sheet.cell(header_row, column).value
                if not protocol:
                    continue
                if (
                    sheet.cell(method_row, column).value is not None
                    and sheet.cell(method_row, column + 1).value is not None
                ):
                    continue
                try:
                    config_path = protocol_to_config[protocol]
                except KeyError as error:
                    raise ValueError(f"No config for workbook protocol {protocol}") from error
                jobs.append((seed, config_path, protocol))
    workbook.close()

    for seed, config_path, protocol in jobs:
        print(f"{seed}\t{config_path}\t{protocol}")


if __name__ == "__main__":
    main()
