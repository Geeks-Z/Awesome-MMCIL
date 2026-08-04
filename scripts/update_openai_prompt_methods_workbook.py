#!/usr/bin/env python3
"""Update the seed-1993 OpenAI CLIP workbook from prompt-method queue logs."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenAI_CLIP_ViT-B16_Baselines.xlsx"
SHEET = "Seed1993"
SEED = 1993
METHOD_LABELS = {
    "l2p": "L2P",
    "dualprompt": "DualPrompt",
    "coda": "CODA-Prompt",
    "proof": "PROOF",
}
DATASET_ALIASES = {
    "aircraft": "aircraft",
    "cifar100": "cifar224",
    "cifar224": "cifar224",
    "cars": "cars",
    "inr": "imagenetr",
    "imagenetr": "imagenetr",
    "cub": "cub",
    "ucf": "ucf101",
    "ucf101": "ucf101",
    "sun": "sun",
    "food": "food101",
    "food101": "food101",
    "objectnet": "objectnet",
}
QUEUE_LOGS = [
    ROOT / "results" / "openai_clip_vit_b16_seed1993" / name
    for name in (
        "l2p_gpu1_tmux.out",
        "dualprompt_gpu2_tmux.out",
        "coda_prompt_gpu5_tmux.out",
        "proof_shard0_gpu6_tmux.out",
        "proof_shard0_gpu6_retry.out",
        "proof_shard1_gpu7_tmux.out",
        "l2p_gpu1_retry_fd65535.out",
        "l2p_food_b0_gpu1.out",
        "l2p_food_b50_gpu5.out",
        "l2p_sun_b0_gpu6.out",
        "l2p_sun_b150_gpu7.out",
    )
]
START = re.compile(
    r"^==== START .*?: (?P<config>configs/[^,]+), seed=(?P<seed>\d+), GPU=\d+ ====$",
    re.MULTILINE,
)
FINISHED = re.compile(
    r"Finished\s+(?P<dataset>\S+)_init(?P<init>\d+)_inc(?P<increment>\d+)\s+seed=(?P<seed>\d+)"
)
AVERAGE = re.compile(r"Average Accuracy \(Top1\):\s*(?P<value>[0-9.]+)")
LAST = re.compile(r"Last Accuracy:\s*(?P<value>[0-9.]+)")
HEADER = re.compile(r"(.+?) B(\d+) Inc(\d+)")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--sheet", default=SHEET)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--allow-incomplete", action="store_true")
    return parser.parse_args()


def normalise_dataset(value: str) -> str:
    key = value.lower().replace("-", "").replace("_", "")
    return DATASET_ALIASES[key]


def config_protocol(config_path: str):
    path = ROOT / config_path
    with path.open() as config_file:
        config = json.load(config_file)
    method = str(config["model_name"]).lower()
    return (
        METHOD_LABELS[method],
        normalise_dataset(str(config["dataset"])),
        int(config["init_cls"]),
        int(config["increment"]),
    )


def parse_log(path: Path):
    content = path.read_text(errors="replace")
    starts = list(START.finditer(content))
    parsed = {}
    for index, start in enumerate(starts):
        if int(start.group("seed")) != SEED:
            continue
        end = starts[index + 1].start() if index + 1 < len(starts) else len(content)
        block = content[start.start() : end]
        finished = FINISHED.search(block)
        averages = AVERAGE.findall(block)
        lasts = LAST.findall(block)
        if not (finished and averages and lasts):
            continue
        key = config_protocol(start.group("config"))
        terminal = (
            normalise_dataset(finished.group("dataset")),
            int(finished.group("init")),
            int(finished.group("increment")),
        )
        if terminal != key[1:] or int(finished.group("seed")) != SEED:
            raise ValueError(f"{path}: terminal output does not match {start.group('config')}")
        metrics = (round(float(averages[-1]), 2), round(float(lasts[-1]), 2))
        if key in parsed and parsed[key] != metrics:
            raise ValueError(f"{path}: conflicting duplicate result for {key}")
        parsed[key] = metrics
    return parsed


def expected_protocols():
    expected = set()
    for directory in ("l2p", "dualprompt", "coda_prompt", "proof"):
        for path in (ROOT / "configs" / directory).glob("*.json"):
            with path.open() as config_file:
                config = json.load(config_file)
            method = str(config["model_name"]).lower()
            if method not in METHOD_LABELS:
                continue
            expected.add(
                (
                    METHOD_LABELS[method],
                    normalise_dataset(str(config["dataset"])),
                    int(config["init_cls"]),
                    int(config["increment"]),
                )
            )
    return expected


def workbook_cells(ws):
    cells = {}
    for row in range(1, ws.max_row + 1):
        method = ws.cell(row, 1).value
        if method not in METHOD_LABELS.values():
            continue
        header_row = max(header for header in range(1, row) if ws.cell(header, 1).value == "Method")
        for col in range(4, ws.max_column + 1, 2):
            header = ws.cell(header_row, col).value
            if not isinstance(header, str):
                continue
            match = HEADER.fullmatch(header.strip())
            if not match:
                continue
            dataset, base, increment = match.groups()
            init = int(increment) if int(base) == 0 else int(base)
            cells[(method, normalise_dataset(dataset), init, int(increment))] = (row, col)
    return cells


def main():
    args = parse_args()
    results = {}
    for path in QUEUE_LOGS:
        if not path.is_file():
            continue
        for key, metrics in parse_log(path).items():
            if key in results and results[key] != metrics:
                raise ValueError(f"Conflicting results across queue logs for {key}")
            results[key] = metrics

    expected = expected_protocols()
    missing = expected.difference(results)
    if missing and not args.allow_incomplete:
        raise ValueError(f"Missing {len(missing)} protocols: {sorted(missing)}")

    workbook = load_workbook(args.workbook)
    ws = workbook[args.sheet]
    cells = workbook_cells(ws)
    if set(results).difference(cells):
        raise ValueError(f"No workbook cells for {sorted(set(results).difference(cells))}")

    changes = []
    for key, (average, last) in sorted(results.items()):
        row, column = cells[key]
        for target_column, value in ((column, average), (column + 1, last)):
            cell = ws.cell(row, target_column)
            if cell.value != value:
                changes.append(f"{ws.title}!{cell.coordinate}: {cell.value!r} -> {value!r}")
                cell.value = value

    print(f"Parsed {len(results)}/{len(expected)} protocols; {'would write' if args.dry_run else 'writing'} {len(changes)} cells")
    for change in changes:
        print(change)
    if missing:
        print(f"Pending {len(missing)} protocols: {sorted(missing)}")
    if not args.dry_run:
        workbook.save(args.workbook)
        print(f"Updated: {args.workbook}")


if __name__ == "__main__":
    main()
