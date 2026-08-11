#!/usr/bin/env python3
"""Fill completed RanPAC/FeCAM results into the CLIP baseline workbooks."""

import argparse
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

RUNS = (
    {
        "name": "OpenAI CLIP",
        "workbook": ROOT / "MMCL_OpenAI_CLIP_ViT-B16_Baselines.xlsx",
        "marker_root": ROOT
        / "logs/OpenAI_CLIP_ViTB16/queue/RanPAC-FeCAM-seed1993/done",
        "log_root": ROOT / "logs/OpenAI_CLIP_ViTB16",
        "backbone_token": "openai_clip",
        "expected_seeds": {1993},
    },
    {
        "name": "OpenCLIP LAION400M",
        "workbook": ROOT / "MMCL_OpenCLIP_ViT-B16_LAION-400M_Baselines.xlsx",
        "marker_root": ROOT
        / "logs/OpenCLIP_LAION400M_ViTB16/queue/RanPAC-FeCAM/done",
        "log_root": ROOT / "logs/OpenCLIP_LAION400M_ViTB16",
        "backbone_token": "clip",
        "expected_seeds": {0, 42, 1993},
    },
)

METHOD_DIR = {"ranpac": "RanPAC", "fecam": "FeCAM"}
DATASET_LABEL = {
    "aircraft": "Aircraft",
    "cars": "Cars",
    "cifar224": "CIFAR100",
    "cub": "CUB",
    "food101": "Food",
    "imagenetr": "INR",
    "objectnet": "ObjectNet",
    "sun": "SUN",
    "ucf101": "UCF",
}

AVERAGE_RE = re.compile(r"Average Accuracy \(CNN top1\):\s*([0-9.]+)")
CURVE_RE = re.compile(r"CNN top1 curve:\s*\[([^\]]+)\]")
SEED_RE = re.compile(r"_seed(\d+)\.done$")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--write", action="store_true", help="Persist updates; default is dry-run."
    )
    return parser.parse_args()


def parse_marker(marker):
    values = {}
    for line in marker.read_text().splitlines():
        if "=" in line:
            key, value = line.split("=", 1)
            values[key] = value
    match = SEED_RE.search(marker.name)
    if not match:
        raise ValueError(f"Cannot parse seed from {marker}")
    values["seed"] = int(match.group(1))
    if "config" not in values:
        raise ValueError(f"Completion marker lacks config path: {marker}")
    return values


def parse_result(log_path):
    text = log_path.read_text(errors="replace")
    averages = AVERAGE_RE.findall(text)
    curves = CURVE_RE.findall(text)
    if not averages or not curves:
        raise ValueError(f"Final metrics are missing from {log_path}")
    curve = [float(value.strip()) for value in curves[-1].split(",")]
    return round(float(averages[-1]), 2), round(curve[-1], 2), len(curve)


def result_from_marker(run, marker):
    marker_data = parse_marker(marker)
    seed = marker_data["seed"]
    if seed not in run["expected_seeds"]:
        raise ValueError(f"Unexpected seed {seed} in {marker}")

    config_path = Path(marker_data["config"])
    if not config_path.is_absolute():
        config_path = ROOT / config_path
    config = json.loads(config_path.read_text())
    method_key = str(config["model_name"]).lower()
    method = METHOD_DIR[method_key]
    dataset = str(config["dataset"])
    init_cls = int(config["init_cls"])
    increment = int(config["increment"])
    # The repository encodes the B0 protocol as init_cls == increment, while
    # trainer.py normalizes that value back to 0 when naming result logs.
    displayed_init_cls = 0 if init_cls == increment else init_cls
    protocol = f"{DATASET_LABEL[dataset]} B{displayed_init_cls} Inc{increment}"
    log_path = (
        run["log_root"]
        / method
        / (
            f"{dataset}_{run['backbone_token']}_"
            f"{displayed_init_cls}_{increment}_{seed}.log"
        )
    )
    if not log_path.is_file():
        raise FileNotFoundError(log_path)
    average, last, curve_length = parse_result(log_path)
    return {
        "method": method,
        "seed": seed,
        "protocol": protocol,
        "average": average,
        "last": last,
        "curve_length": curve_length,
        "log_path": log_path,
    }


def locate_cells(sheet, protocol, method):
    header = None
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == protocol:
                header = cell
                break
        if header is not None:
            break
    if header is None:
        raise ValueError(f"Protocol {protocol!r} is absent from {sheet.title}")

    next_header_row = sheet.max_row + 1
    for row_index in range(header.row + 1, sheet.max_row + 1):
        if sheet.cell(row_index, 1).value == "Method":
            next_header_row = row_index
            break
    method_row = None
    for row_index in range(header.row + 1, next_header_row):
        if sheet.cell(row_index, 1).value == method:
            method_row = row_index
            break
    if method_row is None:
        raise ValueError(
            f"Method {method!r} is absent from the {protocol!r} block in {sheet.title}"
        )
    return sheet.cell(method_row, header.column), sheet.cell(
        method_row, header.column + 1
    )


def update_workbook(run, write):
    markers = sorted(run["marker_root"].glob("*.done"))
    results = [result_from_marker(run, marker) for marker in markers]
    workbook = load_workbook(run["workbook"])
    changes = []

    for result in results:
        sheet_name = f"Seed{result['seed']}"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet {sheet_name} in {run['workbook']}")
        sheet = workbook[sheet_name]
        average_cell, last_cell = locate_cells(
            sheet, result["protocol"], result["method"]
        )
        for cell, metric_name, value in (
            (average_cell, "A_bar", result["average"]),
            (last_cell, "A_B", result["last"]),
        ):
            old_value = cell.value
            if old_value is not None and float(old_value) != value:
                raise ValueError(
                    f"Refusing to overwrite {run['workbook'].name}:{sheet_name}!"
                    f"{cell.coordinate}: existing={old_value}, parsed={value}"
                )
            if old_value is None:
                cell.value = value
                changes.append(
                    (
                        sheet_name,
                        cell.coordinate,
                        result["method"],
                        result["protocol"],
                        metric_name,
                        value,
                    )
                )

    expected_result_count = len(markers)
    if len(results) != expected_result_count:
        raise AssertionError("Not every completion marker produced a result")

    print(
        f"{run['name']}: markers={len(markers)}, results={len(results)}, "
        f"new_cells={len(changes)}"
    )
    for change in changes:
        print("  {}!{} {} {} {}={}".format(*change))

    if write and changes:
        temporary = run["workbook"].with_name(
            run["workbook"].stem + ".tmp" + run["workbook"].suffix
        )
        workbook.save(temporary)
        os.replace(temporary, run["workbook"])
        print(f"  saved={run['workbook']}")
    workbook.close()
    return len(results), len(changes)


def main():
    args = parse_args()
    totals = [update_workbook(run, args.write) for run in RUNS]
    print(
        "TOTAL results={} new_cells={} mode={}".format(
            sum(item[0] for item in totals),
            sum(item[1] for item in totals),
            "write" if args.write else "dry-run",
        )
    )


if __name__ == "__main__":
    main()
