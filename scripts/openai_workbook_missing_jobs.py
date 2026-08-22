#!/usr/bin/env python3
"""Print blank OpenAI CLIP workbook jobs for native MMCL methods."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenAI_CLIP_ViT-B16_Baselines.xlsx"
# key: (workbook label, configuration directory, model_name in configuration)
METHOD_SPECS = {
    "zs_clip": ("ZS-CLIP", "zs_clip", "zs_clip"),
    "l2p": ("L2P", "l2p", "l2p"),
    "dualprompt": ("DualPrompt", "dualprompt", "dualprompt"),
    "coda": ("CODA-Prompt", "coda_prompt", "coda"),
    "ranpac": ("RanPAC", "ranpac", "ranpac"),
    "fecam": ("FeCAM", "fecam", "fecam"),
    "simplecil": ("SimpleCIL", "simplecil", "simplecil"),
    "rapf": ("RAPF", "rapf", "rapf"),
    "proof": ("PROOF", "proof", "proof"),
    "engine": ("ENGINE", "engine", "engine"),
    "clg_cbm": ("CLG-CBM", "CLG-CBM", "CLG-CBM"),
    "bofa": ("BOFA", "bofa", "bofa"),
    "area": ("AREA", "area", "area"),
}
DATASET_LABELS = {
    "aircraft": "Aircraft", "cars": "Cars", "cifar224": "CIFAR100",
    "cub": "CUB", "food101": "Food", "imagenetr": "INR",
    "objectnet": "ObjectNet", "sun": "SUN", "ucf101": "UCF",
}


def protocol_for(config: dict) -> str:
    init_cls, increment = int(config["init_cls"]), int(config["increment"])
    base = 0 if init_cls == increment else init_cls
    return f"{DATASET_LABELS[config['dataset']]} B{base} Inc{increment}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--method",
        action="append",
        choices=sorted(METHOD_SPECS),
        help="Native method key; repeat to select a subset (default: all).",
    )
    args = parser.parse_args()
    selected = args.method or sorted(METHOD_SPECS)
    configs_by_method = {}
    for method in selected:
        _, directory, model_name = METHOD_SPECS[method]
        protocol_to_config = {}
        for config_path in sorted((ROOT / "configs" / directory).glob("*.json")):
            # Recovery snapshots encode a seed/GPU suffix and are not canonical
            # benchmark configurations.
            if "_seed" in config_path.stem:
                continue
            config = json.loads(config_path.read_text())
            if str(config["model_name"]) != model_name:
                raise ValueError(f"Unexpected model_name in {config_path}")
            protocol = protocol_for(config)
            if protocol in protocol_to_config:
                raise ValueError(f"Duplicate protocol config for {method}: {protocol}")
            protocol_to_config[protocol] = config_path
        if len(protocol_to_config) != 18:
            raise ValueError(f"Expected 18 configurations for {method}, found {len(protocol_to_config)}")
        configs_by_method[method] = protocol_to_config

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        if not sheet.title.startswith("Seed"):
            continue
        headers = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Method"]
        for index, header in enumerate(headers):
            end = headers[index + 1] if index + 1 < len(headers) else sheet.max_row + 1
            for method in selected:
                method_name, _, _ = METHOD_SPECS[method]
                method_row = next(
                    (row for row in range(header + 1, end) if sheet.cell(row, 1).value == method_name),
                    None,
                )
                if method_row is None:
                    raise ValueError(f"{method_name} not found in {sheet.title}")
                for column in range(4, sheet.max_column + 1, 2):
                    protocol = sheet.cell(header, column).value
                    if not protocol:
                        continue
                    left, right = sheet.cell(method_row, column).value, sheet.cell(method_row, column + 1).value
                    if left is not None and right is not None:
                        continue
                    config = configs_by_method[method].get(protocol)
                    if config is None:
                        raise ValueError(f"No {method} configuration for {protocol}")
                    print(f"{method}\t{sheet.title.removeprefix('Seed')}\t{config}\t{protocol}")
    workbook.close()


if __name__ == "__main__":
    main()
