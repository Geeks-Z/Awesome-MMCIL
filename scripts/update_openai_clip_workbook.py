#!/usr/bin/env python3
"""Write validated OpenAI CLIP ViT-B/16 seed-1993 results to the workbook.

The launch queues contain different methods in one stdout file, so this parser
uses each queue's ``START`` marker to keep model results separate.  It only
writes a protocol after finding its terminal average and last accuracies.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_Baselines.xlsx"
SHEET = "OpenAI CLIP ViT-B16"
SEED = 1993

METHOD_LABELS = {
    "engine": "ENGINE",
    "rapf": "RAPF",
    "clg-cbm": "CLG-CBM",
    "bofa": "BOFA",
    "zs_clip": "ZS-CLIP",
    "simplecil": "SimpleCIL",
}

DATASET_ALIASES = {
    "aircraft": "aircraft",
    "cifar100": "cifar224",
    "cifar224": "cifar224",
    "cars": "cars",
    "inr": "imagenetr",
    "imagenetr": "imagenetr",
    "cub": "cub",
    "ucf": "ucf101",
    "ucf101": "ucf101",
    "sun": "sun",
    "food": "food101",
    "food101": "food101",
    "objectnet": "objectnet",
}

QUEUE_LOGS = [
    ROOT / "results" / "ENGINE-OpenAI-CLIP-SEED_1993-GPU1-retry1.out",
    ROOT / "results" / "ENGINE-OpenAI-CLIP-SEED_1993-GPU2-retry1.out",
    *[
        ROOT / "results" / f"RAPF-CLG-CBM-BOFA-OpenAI-CLIP-SEED_1993-GPU{gpu}.out"
        for gpu in range(4)
    ],
]

START = re.compile(
    r"^==== START (?P<launcher_method>ENGINE|rapf|clg_cbm|bofa|zs_clip|simplecil) "
    r"OpenAI ViT-B/16: (?P<config>[^,]+), seed=(?P<seed>\d+), GPU=\d+ ====$",
    re.MULTILINE,
)
FINISHED = re.compile(
    r"Finished\s+(?P<dataset>\S+)_init(?P<init>\d+)_inc(?P<increment>\d+) "
    r"seed=(?P<seed>\d+)"
)
AVERAGE = re.compile(r"Average Accuracy \(Top1\):\s*(?P<value>[0-9.]+)")
LAST = re.compile(r"Last Accuracy:\s*(?P<value>[0-9.]+)")
PROTOCOL_HEADER = re.compile(r"(.+?) B(\d+) Inc(\d+)")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--sheet", default=SHEET)
    parser.add_argument(
        "--methods",
        nargs="+",
        choices=sorted(METHOD_LABELS),
        default=["engine", "rapf", "clg-cbm", "bofa"],
    )
    parser.add_argument(
        "--queue-log",
        type=Path,
        action="append",
        help="Result queue log to parse; repeat for each input log.",
    )
    return parser.parse_args()


def normalise_dataset(value: str) -> str:
    key = value.lower().replace("-", "").replace("_", "")
    if key not in DATASET_ALIASES:
        raise ValueError(f"Unsupported dataset {value!r}")
    return DATASET_ALIASES[key]


def protocol_from_config(config_name: str) -> tuple[str, str, int, int]:
    """Return the canonical method/dataset/init/increment from a JSON config."""
    matches = list(ROOT.glob(f"configs/*/{config_name}"))
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one config for {config_name!r}, found {matches}")
    with matches[0].open() as config_file:
        config = json.load(config_file)
    method = str(config["model_name"]).lower()
    if method not in METHOD_LABELS:
        raise ValueError(f"{config_name}: unsupported method {method!r}")
    return (
        METHOD_LABELS[method],
        normalise_dataset(str(config["dataset"])),
        int(config["init_cls"]),
        int(config["increment"]),
    )


def expected_protocols(methods: set[str]) -> set[tuple[str, str, int, int]]:
    expected: set[tuple[str, str, int, int]] = set()
    for directory in ("engine", "rapf", "CLG-CBM", "bofa", "zs_clip", "simplecil"):
        for path in (ROOT / "configs" / directory).glob("*.json"):
            with path.open() as config_file:
                config = json.load(config_file)
            method = str(config["model_name"]).lower()
            if method not in methods or SEED not in config.get("seed", []):
                continue
            expected.add(
                (
                    METHOD_LABELS[method],
                    normalise_dataset(str(config["dataset"])),
                    int(config["init_cls"]),
                    int(config["increment"]),
                )
            )
    return expected


def parse_queue(path: Path) -> dict[tuple[str, str, int, int], tuple[float, float]]:
    content = path.read_text(errors="replace")
    starts = list(START.finditer(content))
    if not starts:
        raise ValueError(f"{path.name}: no OpenAI queue start markers")

    parsed: dict[tuple[str, str, int, int], tuple[float, float]] = {}
    for index, start in enumerate(starts):
        if int(start.group("seed")) != SEED:
            continue
        end = starts[index + 1].start() if index + 1 < len(starts) else len(content)
        block = content[start.start() : end]
        finished = FINISHED.search(block)
        averages = AVERAGE.findall(block)
        lasts = LAST.findall(block)
        if not (finished and averages and lasts):
            raise ValueError(f"{path.name}: incomplete result for {start.group('config')}")
        if int(finished.group("seed")) != SEED:
            raise ValueError(f"{path.name}: mismatched terminal seed for {start.group('config')}")

        method, dataset, init, increment = protocol_from_config(start.group("config"))
        finished_protocol = (
            normalise_dataset(finished.group("dataset")),
            int(finished.group("init")),
            int(finished.group("increment")),
        )
        if finished_protocol != (dataset, init, increment):
            raise ValueError(
                f"{path.name}: terminal protocol {finished_protocol} does not match "
                f"{start.group('config')}"
            )
        launcher_label = start.group("launcher_method").replace("_", "-")
        if launcher_label.upper() != method.upper():
            raise ValueError(f"{path.name}: launcher/config method mismatch for {start.group('config')}")

        key = (method, dataset, init, increment)
        if key in parsed:
            raise ValueError(f"{path.name}: duplicate result for {key}")
        parsed[key] = (round(float(averages[-1]), 2), round(float(lasts[-1]), 2))
    return parsed


def protocol_columns(ws, header_row: int) -> list[tuple[int, tuple[str, int, int]]]:
    columns = []
    for column in range(1, ws.max_column + 1):
        value = ws.cell(header_row, column).value
        if not isinstance(value, str):
            continue
        match = PROTOCOL_HEADER.fullmatch(value.strip())
        if not match:
            continue
        dataset_label, base, increment = match.groups()
        init = int(increment) if int(base) == 0 else int(base)
        columns.append((column, (normalise_dataset(dataset_label), init, int(increment))))
    if not columns:
        raise ValueError(f"{ws.title}!{header_row}: no protocol columns")
    return columns


def write_results(
    results: dict[tuple[str, str, int, int], tuple[float, float]],
    dry_run: bool,
    workbook_path: Path,
    sheet_name: str,
):
    workbook = load_workbook(workbook_path)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook has no {sheet_name!r} sheet")
    ws = workbook[sheet_name]
    header_rows = [row for row in range(1, ws.max_row + 1) if ws.cell(row, 1).value == "Method"]
    if len(header_rows) != 3:
        raise ValueError(f"{SHEET}: expected three method headers, got {len(header_rows)}")

    changes = []
    written = set()
    for index, header_row in enumerate(header_rows):
        next_header = header_rows[index + 1] if index + 1 < len(header_rows) else ws.max_row + 1
        rows = {
            str(ws.cell(row, 1).value): row
            for row in range(header_row + 1, next_header)
            if isinstance(ws.cell(row, 1).value, str)
        }
        for column, (dataset, init, increment) in protocol_columns(ws, header_row):
            for method, row in rows.items():
                key = (method, dataset, init, increment)
                metrics = results.get(key)
                if metrics is None:
                    continue
                written.add(key)
                for target_column, value in ((column, metrics[0]), (column + 1, metrics[1])):
                    cell = ws.cell(row, target_column)
                    if cell.value != value:
                        changes.append(f"{ws.title}!{cell.coordinate}: {cell.value!r} -> {value!r}")
                        cell.value = value

    unmatched = sorted(set(results).difference(written))
    if unmatched:
        raise ValueError(f"No matching workbook cell for {unmatched}")
    print(f"{'Would write' if dry_run else 'Wrote'} {len(changes)} cells for {len(written)} protocols")
    for change in changes:
        print(change)
    if not dry_run:
        workbook.save(workbook_path)
        print(f"Updated: {workbook_path}")


def main():
    args = parse_args()
    methods = set(args.methods)
    queue_logs = args.queue_log or QUEUE_LOGS
    results: dict[tuple[str, str, int, int], tuple[float, float]] = {}
    for path in queue_logs:
        if not path.is_file():
            raise FileNotFoundError(path)
        parsed = parse_queue(path)
        overlap = set(results).intersection(parsed)
        if overlap:
            raise ValueError(f"Duplicate protocols across queue logs: {sorted(overlap)}")
        results.update(parsed)

    expected = expected_protocols(methods)
    if set(results) != expected:
        raise ValueError(
            f"Expected {len(expected)} completed protocols, found {len(results)}; "
            f"missing={sorted(expected.difference(results))}, "
            f"unexpected={sorted(set(results).difference(expected))}"
        )
    write_results(results, args.dry_run, args.workbook, args.sheet)


if __name__ == "__main__":
    main()
