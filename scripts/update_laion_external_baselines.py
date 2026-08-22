#!/usr/bin/env python3
"""Fill complete LAION-400M PromptFusion and MoE-Adapters log results."""

import argparse
import ast
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenCLIP_ViT-B16_LAION-400M_Baselines.xlsx"
LOG_ROOT = ROOT / "logs" / "OpenCLIP_LAION400M_ViTB16"
DATASET_LABELS = {
    "aircraft": "Aircraft",
    "cars": "Cars",
    "cifar224": "CIFAR100",
    "cub": "CUB",
    "food101": "Food",
    "imagenetr": "INR",
    "objectnet": "ObjectNet",
    "sun": "SUN",
    "ucf101": "UCF",
}
METHODS = ("PromptFusion", "MoE-Adapters")
LOG_NAME_RE = re.compile(
    r"(cifar224|imagenetr|cub)_clip_(0|50|100)_(10|20)_(0|42|1993|2026)\.log$"
)
AVERAGE_RE = re.compile(r"Average Accuracy \(CNN top1\):\s*([0-9.]+)")
CURVE_RE = re.compile(r"CNN top1 curve:\s*(\[[^\n]+\])")


def expected_tasks(base: int) -> int:
    return 10 if base == 0 else 6


def parse_log(method: str, path: Path):
    match = LOG_NAME_RE.fullmatch(path.name)
    if match is None:
        return None
    dataset, base_text, increment_text, seed_text = match.groups()
    base, increment, seed = map(int, (base_text, increment_text, seed_text))
    content = path.read_text(errors="replace")
    averages = AVERAGE_RE.findall(content)
    curves = CURVE_RE.findall(content)
    if not averages or not curves or "Backbone: OpenCLIP ViT-B/16 LAION-400M" not in content:
        return None
    curve = ast.literal_eval(curves[-1])
    if not isinstance(curve, list) or len(curve) != expected_tasks(base):
        return None
    if not all(isinstance(value, (int, float)) for value in curve):
        return None
    protocol = f"{DATASET_LABELS[dataset]} B{base} Inc{increment}"
    return seed, method, protocol, round(float(averages[-1]), 2), round(float(curve[-1]), 2)


def locate_cells(sheet, method: str, protocol: str):
    headers = [
        row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Method"
    ]
    for index, header_row in enumerate(headers):
        next_header = headers[index + 1] if index + 1 < len(headers) else sheet.max_row + 1
        protocol_column = next(
            (
                column
                for column in range(4, sheet.max_column + 1, 2)
                if sheet.cell(header_row, column).value == protocol
            ),
            None,
        )
        if protocol_column is None:
            continue
        for row in range(header_row + 1, next_header):
            if sheet.cell(row, 1).value == method:
                return sheet.cell(row, protocol_column), sheet.cell(row, protocol_column + 1)
    raise ValueError(f"Cannot locate {method} {protocol} in {sheet.title}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()
    results = []
    for method in METHODS:
        for path in sorted((LOG_ROOT / method).glob("*.log")):
            parsed = parse_log(method, path)
            if parsed is not None:
                results.append(parsed)

    workbook = load_workbook(WORKBOOK)
    changes = []
    for seed, method, protocol, average, last in results:
        average_cell, last_cell = locate_cells(workbook[f"Seed{seed}"], method, protocol)
        for cell, value in ((average_cell, average), (last_cell, last)):
            if cell.value is not None and float(cell.value) != value:
                raise ValueError(
                    f"Refusing to overwrite {cell.parent.title}!{cell.coordinate}: "
                    f"existing={cell.value}, parsed={value}"
                )
            if cell.value is None:
                cell.value = value
                changes.append(f"{cell.parent.title}!{cell.coordinate}={value}")

    print(f"results={len(results)} new_cells={len(changes)}")
    for change in changes:
        print(change)
    if args.write and changes:
        workbook.save(WORKBOOK)
        print(f"saved={WORKBOOK}")
    workbook.close()


if __name__ == "__main__":
    main()
