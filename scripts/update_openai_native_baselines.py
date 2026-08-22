#!/usr/bin/env python3
"""Fill blank OpenAI CLIP workbook cells from complete native-method logs."""

import argparse
import ast
import math
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "MMCL_OpenAI_CLIP_ViT-B16_Baselines.xlsx"
OPENAI_LOG_ROOT = ROOT / "logs" / "OpenAI_CLIP_ViTB16"
METHODS = (
    "ZS-CLIP",
    "L2P",
    "DualPrompt",
    "CODA-Prompt",
    "RanPAC",
    "FeCAM",
    "SimpleCIL",
    "RAPF",
    "PROOF",
    "ENGINE",
    "CLG-CBM",
    "BOFA",
    "AREA",
)
DATASET_LABELS = {
    "aircraft": "Aircraft", "cars": "Cars", "cifar224": "CIFAR100",
    "cub": "CUB", "food101": "Food", "imagenetr": "INR",
    "objectnet": "ObjectNet", "sun": "SUN", "ucf101": "UCF",
}
TOTAL_CLASSES = {
    "aircraft": 100, "cars": 100, "cifar224": 100, "cub": 200,
    "food101": 100, "imagenetr": 200, "objectnet": 200, "sun": 300,
    "ucf101": 100,
}
LOG_PATTERN = re.compile(
    r"(aircraft|cars|cifar224|cub|food101|imagenetr|objectnet|sun|ucf101)_"
    r"openai_clip_(0|50|100|150)_(10|20|30)_(0|42|1993|2026)\.log$"
)
AVERAGE = re.compile(r"Average Accuracy \(CNN top1\):\s*([0-9.]+)")
CURVE = re.compile(r"CNN top1 curve:\s*(\[[^\n]+\])")


def close_pair(left, right) -> bool:
    return abs(left[0] - right[0]) <= 0.0051 and abs(left[1] - right[1]) <= 0.0051


def locate_cells(sheet, method: str, dataset: str, base: int, increment: int):
    protocol = f"{DATASET_LABELS[dataset]} B{base} Inc{increment}"
    headers = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Method"]
    for index, header in enumerate(headers):
        end = headers[index + 1] if index + 1 < len(headers) else sheet.max_row + 1
        column = next((col for col in range(4, sheet.max_column + 1, 2) if sheet.cell(header, col).value == protocol), None)
        if column is None:
            continue
        row = next((row for row in range(header + 1, end) if sheet.cell(row, 1).value == method), None)
        if row is not None:
            return row, column
    raise ValueError(f"No workbook cell for {method} {protocol} in {sheet.title}")


def complete_results():
    results = {}
    for method in METHODS:
        for path in sorted((OPENAI_LOG_ROOT / method).glob("*.log")):
            match = LOG_PATTERN.fullmatch(path.name)
            if match is None:
                continue
            dataset, base_text, increment_text, seed_text = match.groups()
            base, increment, seed = map(int, (base_text, increment_text, seed_text))
            content = path.read_text(errors="replace")
            averages, curves = AVERAGE.findall(content), CURVE.findall(content)
            if not averages or not curves:
                continue
            try:
                curve = ast.literal_eval(curves[-1])
            except (SyntaxError, ValueError):
                continue
            init_cls = increment if base == 0 else base
            expected_tasks = 1 + math.ceil((TOTAL_CLASSES[dataset] - init_cls) / increment)
            if not isinstance(curve, list) or len(curve) != expected_tasks or not all(isinstance(value, (int, float)) for value in curve):
                continue
            key = (seed, method, dataset, base, increment)
            if key in results:
                raise ValueError(f"Duplicate complete log for {key}")
            results[key] = (float(averages[-1]), float(curve[-1]))
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()
    results = complete_results()
    workbook = load_workbook(WORKBOOK)
    changes = []
    for (seed, method, dataset, base, increment), metrics in sorted(results.items()):
        worksheet = workbook[f"Seed{seed}"]
        row, column = locate_cells(worksheet, method, dataset, base, increment)
        target = (round(metrics[0], 2), round(metrics[1], 2))
        existing = (worksheet.cell(row, column).value, worksheet.cell(row, column + 1).value)
        if existing == (None, None):
            worksheet.cell(row, column).value, worksheet.cell(row, column + 1).value = target
            changes.append(f"{worksheet.title}!{worksheet.cell(row, column).coordinate}={target[0]} {worksheet.cell(row, column + 1).coordinate}={target[1]}")
        elif not all(isinstance(value, (int, float)) for value in existing) or not close_pair((float(existing[0]), float(existing[1])), metrics):
            raise ValueError(f"Refusing to overwrite {worksheet.title} {method} {dataset} B{base} Inc{increment}: {existing} != {target}")
    print(f"results={len(results)} new_cells={len(changes)}")
    for change in changes:
        print(change)
    if args.write and changes:
        workbook.save(WORKBOOK)
        print(f"saved={WORKBOOK}")
    workbook.close()


if __name__ == "__main__":
    main()
